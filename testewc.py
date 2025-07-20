from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
doc.add_heading('Relatório de vendas', level=1)
doc.add_paragraph('Paragrafo').add_run('Texto em negrito').bold = True
doc.add_paragraph('Teste').add_run('Texto em itálico').italic = True
doc.add_heading('Título Nível 1', level=1)
doc.add_heading('Título Nível 2', level=2)

p = doc.add_paragraph('''
Este é um exemplo de texto totalmente justificado. O alinhamento justificado é comum em relatórios, livros e documentos formais, pois proporciona uma aparência organizada e profissional ao texto. Todas as linhas são alinhadas tanto à margem esquerda quanto à margem direita, ajustando o espaçamento entre as palavras.
''')

# Define alinhamento como JUSTIFY
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# Salva o documento
doc.save('texto_justificado.docx')